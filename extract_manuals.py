import os
import sys
import json
from pathlib import Path

try:
    from PyPDF2 import PdfReader
except ImportError:
    PdfReader = None

try:
    import docx
except ImportError:
    docx = None

try:
    import zipfile
except ImportError:
    zipfile = None

try:
    import rarfile
except ImportError:
    rarfile = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

ROOT = Path(r'C:\Users\DEVICES\Desktop\ADS-B API')
OUTPUT_DIR = ROOT / 'adsb-api' / 'docs' / 'manuals_extracted'
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

def write_text(file_path: Path, text: str):
    out_path = OUTPUT_DIR / (file_path.stem + '.txt')
    out_path.write_text(text, encoding='utf-8')
    print(f'  Extracted -> {out_path.name}  ({len(text)} chars)')

def extract_pdf(path: Path):
    if not PdfReader:
        print(f'  SKIP (no PyPDF2): {path.name}')
        return
    try:
        reader = PdfReader(str(path))
        text = []
        for page in reader.pages:
            text.append(page.extract_text() or '')
        write_text(path, '\n'.join(text))
    except Exception as e:
        print(f'  ERROR extracting PDF {path.name}: {e}')

def extract_docx(path: Path):
    if not docx:
        print(f'  SKIP (no python-docx): {path.name}')
        return
    try:
        doc = docx.Document(str(path))
        text = '\n'.join(p.text for p in doc.paragraphs)
        write_text(path, text)
    except Exception as e:
        print(f'  ERROR extracting DOCX {path.name}: {e}')

def extract_xlsx(path: Path):
    if not openpyxl:
        print(f'  SKIP (no openpyxl): {path.name}')
        return
    try:
        wb = openpyxl.load_workbook(str(path), data_only=True)
        lines = []
        for sheet in wb.sheetnames:
            lines.append(f'=== Sheet: {sheet} ===')
            ws = wb[sheet]
            for row in ws.iter_rows(values_only=True):
                lines.append('\t'.join('' if v is None else str(v) for v in row))
        write_text(path, '\n'.join(lines))
    except Exception as e:
        print(f'  ERROR extracting XLSX {path.name}: {e}')

def extract_rar(path: Path):
    if not rarfile:
        print(f'  SKIP (no rarfile): {path.name}')
        return
    try:
        with rarfile.RarFile(str(path)) as rf:
            for info in rf.infolist():
                if info.is_dir():
                    continue
                ext = Path(info.filename).suffix.lower()
                if ext in ['.pdf', '.docx', '.xlsx']:
                    data = rf.read(info)
                    temp = OUTPUT_DIR / ('_rar_' + Path(info.filename).name)
                    temp.write_bytes(data)
                    print(f'  From RAR: {info.filename}')
                    if ext == '.pdf':
                        extract_pdf(temp)
                    elif ext == '.docx':
                        extract_docx(temp)
                    elif ext == '.xlsx':
                        extract_xlsx(temp)
                    temp.unlink(missing_ok=True)
    except Exception as e:
        print(f'  ERROR extracting RAR {path.name}: {e}')

def extract_zip(path: Path):
    try:
        import zipfile as zf
        with zf.ZipFile(str(path), 'r') as z:
            for info in z.infolist():
                if info.is_dir():
                    continue
                ext = Path(info.filename).suffix.lower()
                if ext in ['.pdf', '.docx', '.xlsx']:
                    data = z.read(info)
                    temp = OUTPUT_DIR / ('_zip_' + Path(info.filename).name)
                    temp.write_bytes(data)
                    print(f'  From ZIP: {info.filename}')
                    if ext == '.pdf':
                        extract_pdf(temp)
                    elif ext == '.docx':
                        extract_docx(temp)
                    elif ext == '.xlsx':
                        extract_xlsx(temp)
                    temp.unlink(missing_ok=True)
    except Exception as e:
        print(f'  ERROR extracting ZIP {path.name}: {e}')

def main():
    # Fix encoding for Chinese filenames on Windows
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')

    # Scan root folder for manuals (not recursing into node_modules etc.)
    exts = {'.pdf', '.docx', '.xlsx', '.rar', '.zip'}
    files = [f for f in ROOT.iterdir() if f.is_file() and f.suffix.lower() in exts]

    print(f'Found {len(files)} manual files in {ROOT}')
    print('=' * 60)

    for path in sorted(files):
        ext = path.suffix.lower()
        print(f'\nProcessing: {path.name}')
        if ext == '.pdf':
            extract_pdf(path)
        elif ext == '.docx':
            extract_docx(path)
        elif ext == '.xlsx':
            extract_xlsx(path)
        elif ext == '.rar':
            extract_rar(path)
        elif ext == '.zip':
            extract_zip(path)

    print('\n' + '=' * 60)
    extracted = list(OUTPUT_DIR.glob('*.txt'))
    print(f'Done! {len(extracted)} text files in {OUTPUT_DIR}')
    for f in sorted(extracted):
        print(f'  {f.name}  ({f.stat().st_size} bytes)')

if __name__ == '__main__':
    main()
